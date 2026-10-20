import docx

doc = docx.Document('Revisi 4 A_Deep_Learning_Architecture_for_Cryptocurrency_Price_Forecasting.docx')

for i, p in enumerate(doc.paragraphs):
    blips = p._p.xpath('.//a:blip')
    if blips:
        for b in blips:
            embed_id = b.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
            target_part = doc.part.related_parts.get(embed_id)
            partname = target_part.partname if target_part else "unknown"
            print(f"Paragraph [{i}] has image: embed={embed_id}, target={partname}, text={p.text[:40]}")

for t_idx, table in enumerate(doc.tables):
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            blips = cell._tc.xpath('.//a:blip')
            if blips:
                for b in blips:
                    embed_id = b.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                    target_part = doc.part.related_parts.get(embed_id)
                    partname = target_part.partname if target_part else "unknown"
                    print(f"Table [{t_idx}] Row [{r_idx}] Cell [{c_idx}] has image: embed={embed_id}, target={partname}, text={cell.text[:40]}")
