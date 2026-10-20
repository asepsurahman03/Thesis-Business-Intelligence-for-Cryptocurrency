"""
Comprehensive revision script to apply all changes to:
Revisi 4 A_Deep_Learning_Architecture_for_Cryptocurrency_Price_Forecasting.docx
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
FIG_DIR = os.path.join(BASE_DIR, "04_Visualizations")
TARGET_DOC = os.path.join(BASE_DIR, "Revisi 4 A_Deep_Learning_Architecture_for_Cryptocurrency_Price_Forecasting.docx")

FIG1_PATH = os.path.join(FIG_DIR, "Fig1_BTC_Price_Volume_300dpi.png")
FIG2_PATH = os.path.join(FIG_DIR, "Fig2_Model_Comparison_300dpi.png")
FIG3_PATH = os.path.join(FIG_DIR, "Fig3_Residual_Distribution_300dpi.png")
FIG4_PATH = os.path.join(FIG_DIR, "Fig4_BI_Dashboard_KPI_300dpi.png")

# Load Document
doc = Document(TARGET_DOC)

def set_para_text(p, text, bold_prefix=None):
    """Sets text of a paragraph cleanly preserving paragraph style."""
    p.text = ""
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.bold = True
    p.add_run(text)

print("Starting in-place revision...")

# -------------------------------------------------------------
# 0. Title & Keywords
# -------------------------------------------------------------
doc.paragraphs[0].text = "A Deep Learning Architecture for Cryptocurrency Price Forecasting Utilizing Business Intelligence"

# 1. Abstract (P[10])
abstract_text = (
    "Forecasting cryptocurrency prices in operational financial settings presents substantial challenges "
    "due to non-stationarity, regime switching, and extreme volatility. This paper presents a four-layer automated "
    "Business Intelligence and Machine Learning (BI-ML) framework that injects sequential predictive models into an "
    "enterprise decision-support system, evaluating out-of-sample forecasts through real-time financial utility metrics "
    "including the Sharpe Ratio and Value-at-Risk (VaR). Four model topologies—Extreme Gradient Boosting (XGBoost), "
    "Long Short-Term Memory (LSTM), Gated Recurrent Unit (GRU), and the Informer Transformer—were rigorously evaluated "
    "under a 24-fold chronological Walk-Forward Validation protocol using 43,786 hourly candlestick observations of "
    "BTC/USDT from March 2021 to March 2026. In terms of point-forecasting accuracy, the GRU achieved the lowest error "
    "(RMSE = 0.1058 ± 0.0034, p < 0.001). However, when evaluated within the operational BI decision engine, a critical "
    "empirical divergence emerged: the LSTM achieved the highest Directional Accuracy (72.3% ± 1.5%) and the highest "
    "annualized Sharpe Ratio (1.45 ± 0.10, with VaR95% = −2.34%), substantially outperforming models with lower point errors. "
    "These findings demonstrate that conventional loss metrics such as RMSE fail to indicate trading efficacy and "
    "risk-adjusted financial utility. Usability evaluation of the enterprise Power BI prototype yielded a System Usability "
    "Scale (SUS) score of 86.4 (SD = 3.2), confirming that domain practitioners can autonomously operate and interpret the "
    "multi-model decision support dashboard for risk-governed portfolio allocation."
)
set_para_text(doc.paragraphs[10], abstract_text, bold_prefix="Abstract—")

# 2. Keywords (P[11])
keywords_text = "Financial Informatics, Business Intelligence, Deep Learning, XGBoost, LSTM, GRU, Informer Transformer, Decision Support System, Risk Management."
set_para_text(doc.paragraphs[11], keywords_text, bold_prefix="Keywords—")

# -------------------------------------------------------------
# 3. Section I: Introduction
# -------------------------------------------------------------
doc.paragraphs[13].text = "I. INTRODUCTION"

intro_p1 = (
    "Cryptocurrency markets exhibit extreme non-stationarity, high-frequency regime shifts, and severe volatility cascades, "
    "frequently invalidating the core assumptions underlying conventional econometric models [1], [2]. Classical time-series "
    "techniques such as Autoregressive Integrated Moving Average (ARIMA) presuppose covariance stationarity and linear "
    "autocorrelation structures—properties routinely violated by intra-day cryptocurrency data [3], [4]. Consequently, deep neural "
    "networks and gradient-boosted tree ensembles have emerged as prominent modeling paradigms due to their capacity to map "
    "complex, nonlinear sequential representations directly from raw multi-feature inputs [5], [6]."
)
set_para_text(doc.paragraphs[14], intro_p1)

intro_p2 = (
    "Despite significant advances in forecasting algorithms, a critical methodological limitation persists in the financial machine "
    "learning literature. Existing studies overwhelmingly benchmark models using statistical error metrics such as Root Mean "
    "Squared Error (RMSE) or Mean Absolute Error (MAE) in isolated laboratory environments. However, for portfolio managers and "
    "risk governance committees, a minimal RMSE does not guarantee commercial viability or risk-adjusted profitability [7], [8]. "
    "Crucial operational questions—such as directional forecasting accuracy, strategy drawdown, and risk-adjusted return under real-world "
    "trading constraints—remain unanswered by error metrics alone [9]. A profound disconnect therefore exists between predictive accuracy "
    "benchmarks and the operational requirements of enterprise Business Intelligence (BI) platforms [10]."
)
set_para_text(doc.paragraphs[15], intro_p2)

intro_p3 = (
    "To bridge this divide, this study develops an end-to-end BI-ML framework that integrates automated multi-model forecasting, "
    "leakage-free chronological Walk-Forward Validation, a high-throughput PostgreSQL Data Mart, and an interactive Microsoft "
    "Power BI decision-support dashboard operating in near real-time. Four distinct algorithmic families are benchmarked across "
    "five years of hourly Bitcoin (BTC/USDT) spot data under identical experimental conditions: XGBoost (tree ensemble), LSTM and "
    "GRU (gated recurrent networks), and the Informer Transformer (probabilistic sparse self-attention). By mapping raw multi-step "
    "point predictions into live trading signals, annualized Sharpe Ratios, and Value-at-Risk (VaR) thresholds, this framework "
    "establishes a rigorous bridge between statistical modeling and enterprise financial decision governance."
)
set_para_text(doc.paragraphs[16], intro_p3)

doc.paragraphs[17].text = "A. Baselines and Deep Sequences (RNNs)"
p18_text = (
    "In financial modeling, tree-based ensemble methods represent the established standard for tabular feature architectures. "
    "Extreme Gradient Boosting (XGBoost) demonstrates superior capacity to mitigate variance through regularized gradient boosting "
    "[11], [12]. Nevertheless, standard decision tree models do not inherently model recurrent temporal transitions across consecutive "
    "time steps (t − 1), a limitation highlighted by Zheng et al. [13] in volatile crypto asset regimes."
)
set_para_text(doc.paragraphs[18], p18_text)

p19_text = (
    "To capture multi-scale temporal dependencies, recurrent neural architectures were introduced. The Long Short-Term Memory (LSTM) "
    "network, formulated by Hochreiter and Schmidhuber [14], regulates gradient flow via input, forget, and output gating mechanisms, "
    "mitigating vanishing gradient phenomena. Subsequently, Cho et al. [15] developed the Gated Recurrent Unit (GRU), which combines "
    "forget and input gates into a streamlined update gate, achieving comparable sequential retention with reduced computational overhead. "
    "Both topologies remain essential benchmarks in quantitative financial engineering [16]–[18]."
)
set_para_text(doc.paragraphs[19], p19_text)

doc.paragraphs[20].text = "B. Transformers in Time Series"
p21_text = (
    "Attention-based Transformer architectures [19] have achieved remarkable success across diverse sequential modeling domains [20]. "
    "For long-sequence time-series forecasting, the Informer Transformer [21] introduced the ProbSparse self-attention mechanism, "
    "reducing time and memory complexity from O(L²) to O(L log L). While newer Transformer derivatives such as PatchTST [22] and "
    "TimesNet [23] continue to emerge, the Informer serves as an established, extensively validated attention baseline for intra-day "
    "cryptocurrency modeling. Importantly, Zerveas et al. [24] observed that attention mechanisms are prone to overfitting on short, "
    "noisy intervals—a phenomenon directly investigated in this study."
)
set_para_text(doc.paragraphs[21], p21_text)

doc.paragraphs[22].text = "C. The Need for BI Integration and Scope Justification"
p23_text = (
    "While comparative algorithmic evaluations are abundant [27], [28], their integration into enterprise Decision Support Systems "
    "(DSS) remains scarce. The primary objective of this work is to bridge this operational gap by coupling deep learning predictions "
    "to an enterprise BI layer that dynamically computes the annualized Sharpe Ratio and Value-at-Risk (VaR) for institutional risk "
    "governance [29]."
)
set_para_text(doc.paragraphs[23], p23_text)

p24_text = (
    "Bitcoin (BTC/USDT) is selected as the primary reference asset for three fundamental reasons. First, BTC constitutes the most liquid, "
    "heavily traded, and structurally influential cryptocurrency, serving as the universal market bellwether [1], [30]. Second, continuous "
    "24/7 trading generates an exceptionally dense sequential signal unobstructed by exchange market-closure artifacts. Third, evaluating "
    "a high-volume Binance spot pair provides a controlled experimental environment that isolates algorithmic performance from the idiosyncratic "
    "liquidity shocks prevalent in lower-cap altcoins."
)
set_para_text(doc.paragraphs[24], p24_text)

# -------------------------------------------------------------
# 4. Section II: Proposed Methodology and System Architecture
# -------------------------------------------------------------
doc.paragraphs[25].text = "II. METHODOLOGY AND SYSTEM ARCHITECTURE"

method_intro = (
    "The proposed end-to-end framework is organized into four modular, sequentially integrated layers. Layer 1 executes automated "
    "high-frequency market data ingestion and verifies series stationarity. Layer 2 performs technical feature extraction, fold-local "
    "normalization, and guarantees temporal leakage prevention. Layer 3 orchestrates model training, hyperparameter optimization, and "
    "parallel out-of-sample inference across the four model families. Finally, Layer 4 translates raw predictive tensors into directional "
    "trading signals, computes portfolio risk metrics (Sharpe Ratio, VaR), and streams the outputs to a live Microsoft Power BI dashboard "
    "via DirectQuery. The end-to-end procedure is formally outlined in Algorithm 1."
)
set_para_text(doc.paragraphs[26], method_intro)

doc.paragraphs[27].text = "A. High-Frequency Market Data Extraction (Layer 1)"
p28_text = (
    "Layer 1 ingests 43,786 consecutive hourly OHLCV candlestick records of BTC/USDT from the Binance Spot API spanning March 2021 to "
    "March 2026 [30]. The dataset represents continuous, verified market execution with 100% temporal continuity, providing an uninterrupted "
    "basis for sequential modeling."
)
set_para_text(doc.paragraphs[28], p28_text)

p29_text = (
    "To ensure mathematical validity for sequential deep learning architectures, all derived feature series were subjected to the "
    "Augmented Dickey-Fuller (ADF) unit root test. In all cases, the null hypothesis of a unit root was decisively rejected at the 5% "
    "significance level (p < 0.05), confirming covariance stationarity across the feature set."
)
set_para_text(doc.paragraphs[29], p29_text)

doc.paragraphs[30].text = "B. Feature Engineering, Volatility Signals, and Temporal Leakage Prevention (Layer 2)"
p31_text = (
    "A comprehensive feature set of 15 technical indicators was engineered to provide multi-scale representations of price momentum, "
    "trend, and market volatility. Moving Average Convergence Divergence (MACD) was computed using standard exponential smoothing [34]:"
)
set_para_text(doc.paragraphs[31], p31_text)

# Note: Equation 2 is at P[45] in original, let's reorganize paragraphs cleanly.
# Let's verify how paragraphs map.

print("Applied section I and initial section II paragraphs.")
