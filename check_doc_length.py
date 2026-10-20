import docx

doc = docx.Document('05_Reports/Manuscript_IJAI_FINAL_V2.docx')
text = "\n".join([p.text for p in doc.paragraphs])
words = text.split()
print("Total words in document:", len(words))

# Find abstract
abstract_text = ""
for p in doc.paragraphs:
    if p.text.startswith('ABSTRACT'):
        continue
    if p.text.startswith('Keywords'):
        break
    if len(abstract_text) > 0 or p.text.strip().startswith('The inherent'):
        abstract_text += p.text + " "

print("Abstract words:", len(abstract_text.split()))
print("Abstract:", abstract_text[:100], "...")
