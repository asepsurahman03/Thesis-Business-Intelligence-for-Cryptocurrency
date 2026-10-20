import docx

doc = docx.Document('05_Reports/Manuscript_IJAI_FINAL_V3.docx')
text = "\n".join([p.text for p in doc.paragraphs])
words = text.split()
print("Total words in V3 document:", len(words))

abstract_text = ""
for p in doc.paragraphs:
    if p.text.startswith('Abstract—'):
        abstract_text = p.text
        break

print("Abstract words:", len(abstract_text.split()))
