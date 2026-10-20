import docx
import re
import os
import win32com.client

input_path = os.path.abspath('05_Reports/Manuscript_IJAI_FINAL_V5.docx')

# 1. Extract all formulas using python-docx
doc = docx.Document(input_path)
full_text = []
for p in doc.paragraphs:
    full_text.append(p.text)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                full_text.append(p.text)

combined_text = '\n'.join(full_text)
# Find all $...$ and $$...$$
pattern = re.compile(r'\$\$.*?\$\$|\$.*?\$', re.DOTALL)
formulas = list(set(pattern.findall(combined_text)))

print(f"Found {len(formulas)} unique formulas.")
for f in formulas[:5]:
    print(" -", f)

# 2. Use win32com to find and convert them
word = win32com.client.Dispatch("Word.Application")
word.Visible = False

doc_com = word.Documents.Open(input_path)

for formula in formulas:
    find = doc_com.Content.Find
    find.ClearFormatting()
    
    # We must handle newlines in formulas if any (usually dotall matches them)
    # Word Find might fail if the formula contains carriage returns.
    # Let's clean the search string or use a careful find.
    search_text = formula
    if len(search_text) > 255:
        print(f"Formula too long for Word Find (max 255 chars): {search_text[:30]}...")
        continue
        
    find.Text = search_text.replace('\n', '^p')
    find.Forward = True
    find.Wrap = 1 # wdFindContinue
    find.Format = False
    find.MatchCase = True
    find.MatchWholeWord = False
    find.MatchWildcards = False
    find.MatchSoundsLike = False
    find.MatchAllWordForms = False

    while find.Execute():
        rng = find.Parent
        
        # Word includes the $ and $$ in the range. 
        # When we add OMath, it might keep the $ signs if we don't remove them.
        # Let's remove the $ signs from the text in the document before converting.
        is_block = formula.startswith('$$')
        clean_text = formula
        if is_block:
            clean_text = formula[2:-2].strip()
        else:
            clean_text = formula[1:-1].strip()
            
        rng.Text = clean_text
        
        # Now rng points to the cleaned text
        doc_com.OMaths.Add(rng)
        
        # After Add, we get an OMath object. We need to build it up.
        # But wait, finding the added OMath can be tricky.
        # Usually it's the last one in the range or we can just iterate OMaths in rng.
        for i in range(1, rng.OMaths.Count + 1):
            rng.OMaths(i).BuildUp()
        
        # Reset the Find range to the whole document after the current position to avoid infinite loops
        # Or just use doc_com.Content.Find again, but since we modified the text, it won't be found again!
        find = doc_com.Content.Find
        find.ClearFormatting()
        find.Text = search_text.replace('\n', '^p')
        find.Forward = True
        find.Wrap = 1

doc_com.Save()
doc_com.Close()
word.Quit()
print("Finished converting formulas natively using Word COM.")
