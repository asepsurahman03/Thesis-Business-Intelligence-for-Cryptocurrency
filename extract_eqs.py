import docx
doc = docx.Document('05_Reports/Manuscript_IJAI_FINAL_V5.docx')

with open('05_Reports/eq_dump.txt', 'w', encoding='utf-8') as f:
    for p in doc.paragraphs:
        if '=' in p.text or '+' in p.text or '-' in p.text or '\\' in p.text or '[' in p.text or '(' in p.text:
            f.write(p.text + '\n')
            
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if '=' in p.text or '+' in p.text or '-' in p.text or '\\' in p.text or '[' in p.text or '(' in p.text:
                        f.write(p.text + '\n')
