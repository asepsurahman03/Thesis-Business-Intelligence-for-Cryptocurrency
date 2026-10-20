import docx

doc = docx.Document('Revisi 4 A_Deep_Learning_Architecture_for_Cryptocurrency_Price_Forecasting.docx')

with open('full_manuscript_review.txt', 'w', encoding='utf-8') as f:
    for i, p in enumerate(doc.paragraphs):
        f.write(f"[{i:03d}] [{p.style.name}] {p.text}\n")
    for t_idx, t in enumerate(doc.tables):
        f.write(f"\n--- TABLE {t_idx} ---\n")
        for r_idx, r in enumerate(t.rows):
            f.write(f"  Row {r_idx}: {' | '.join(c.text.strip().replace(chr(10), ' ') for c in r.cells)}\n")

print("Exported full_manuscript_review.txt")
