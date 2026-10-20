import os
import shutil
from docx import Document

BACKUP_FILE = r'Tesis\026 FORM PERSYARATAN SIDANG TESIS.docx.bak_original'
TARGET_FILE = r'Tesis\026 FORM PERSYARATAN SIDANG TESIS.docx'

if not os.path.exists(BACKUP_FILE):
    shutil.copy2(TARGET_FILE, BACKUP_FILE)

# Start clean from original template
shutil.copy2(BACKUP_FILE, TARGET_FILE)
doc = Document(TARGET_FILE)

# Master Data extracted from Tesis_Asep_Surahman_Sulaeman_2026_rapi.docx
DATA = {
    'nama': 'Asep Surahman Sulaeman',
    'nim': '20250130001',
    'ttl': 'Cianjur, 03 Maret 2003',
    'prodi': 'Magister Informatika',
    'konsentrasi': 'Data Science',
    'alamat': 'Kp. Babakan Kondang, Desa Hegarmanah, Kecamatan Takokak, Kabupaten Cianjur, Jawa Barat',
    'alamat_ortu': 'Kp. Babakan Kondang, Desa Hegarmanah, Kecamatan Takokak, Kabupaten Cianjur, Jawa Barat',
    'hp': '+62 856-5983-8977',
    'hp_ortu': '-',
    'email': 'asep.surahman@nusaputra.ac.id',
    'ketua_pembimbing': 'Prof. Ts. Deden Witarsyah, S.T, Ph.D',
    'nidn_ketua_pembimbing': '0120240060',
    'anggota_pembimbing': 'Assoc. Prof. Dr. Aswan Supriyadi Sunge, S.E., M.Kom',
    'nidn_anggota_pembimbing': '0426018003',
    'kaprodi': 'Assoc. Prof. Dr. Adhi Kusnadi, S.T., M.Si',
    'nidn_kaprodi': '0303037304',
    'judul': 'Kerangka Kerja Business Intelligence untuk Prediksi Cryptocurrency: Studi Komparatif LSTM, GRU, dan Ensemble Baseline Menggunakan Big Data Analytics dengan Integrasi Deep Learning',
    'tanggal': '05 Oktober 2026',
    'semester': 'SEMESTER GANJIL TAHUN AKADEMIK 2026/2027'
}

paras = doc.paragraphs

# --- 1. HALAMAN 1: FORMULIR PENDAFTARAN SIDANG ---
# [0] Header Judul Form
for r in paras[0].runs:
    if '...' in r.text:
        r.text = r.text.replace('...', 'INFORMATIKA')

# [3] Program Studi Header
for r in paras[3].runs:
    if '\t' in r.text:
        r.text = r.text.replace('\t', f': {DATA["prodi"]}\t')

# [6] Sukabumi,
for r in paras[6].runs:
    if 'Sukabumi,' in r.text:
        r.text = f'Sukabumi, {DATA["tanggal"]}'

# [11] Nama Mahasiswa
paras[11].text = DATA['nama']
if paras[11].runs:
    paras[11].runs[0].bold = True

# --- TABEL 0: BIODATA MAHASISWA & PEMBIMBING ---
t0 = doc.tables[0]

def populate_cell(cell, text, bold=False):
    for p in cell.paragraphs:
        for r in p.runs:
            r.text = ''
    if not cell.paragraphs:
        cell.add_paragraph()
    p = cell.paragraphs[0]
    if not p.runs:
        r = p.add_run(text)
    else:
        r = p.runs[0]
        r.text = text
    r.bold = bold

# R0: Nama & NIM
populate_cell(t0.rows[0].cells[1], DATA['nama'])
nim_cell = t0.rows[0].cells[2]
for p in nim_cell.paragraphs:
    for r in p.runs:
        if 'NIM' in r.text:
            r.text = f'NIM: {DATA["nim"]}'
            break
if DATA['nim'] not in nim_cell.text:
    nim_cell.paragraphs[0].text = f'NIM: {DATA["nim"]}'

# R1: TTL
populate_cell(t0.rows[1].cells[1], DATA['ttl'])

# R2: Program Studi
populate_cell(t0.rows[2].cells[1], DATA['prodi'])

# R3: Konsentrasi
populate_cell(t0.rows[3].cells[1], DATA['konsentrasi'])

# R4: Alamat Tinggal
populate_cell(t0.rows[4].cells[1], DATA['alamat'])

# R5: Alamat Tinggal Orang Tua
populate_cell(t0.rows[5].cells[1], DATA['alamat_ortu'])

# R6: No Tlp/HP & HP Ortu
populate_cell(t0.rows[6].cells[1], DATA['hp'])
hp_ortu_cell = t0.rows[6].cells[2]
for p in hp_ortu_cell.paragraphs:
    for r in p.runs:
        if 'No HP Orang Tua' in r.text:
            r.text = f'No HP Orang Tua: {DATA["hp_ortu"]}'
if DATA['hp_ortu'] not in hp_ortu_cell.text:
    hp_ortu_cell.paragraphs[0].text = f'No HP Orang Tua: {DATA["hp_ortu"]}'

# R7: Alamat Email (merged 1-2)
populate_cell(t0.rows[7].cells[1], DATA['email'])

# R8: Ketua Pembimbing (merged 1-2)
populate_cell(t0.rows[8].cells[1], DATA['ketua_pembimbing'])

# R9: Anggota Pembimbing (merged 1-2)
populate_cell(t0.rows[9].cells[1], DATA['anggota_pembimbing'])

# R10: Judul Tesis (merged 1-2)
populate_cell(t0.rows[10].cells[1], DATA['judul'])

# --- 2. HALAMAN 2: LEMBAR PERSYARATAN SIDANG ---
# [14] Nama :
for r in paras[14].runs:
    if ':' in r.text:
        r.text = f': {DATA["nama"]}'

# [15] NIM :
for r in paras[15].runs:
    if ':' in r.text:
        r.text = f': {DATA["nim"]}'

# [16] Program Studi :
for r in paras[16].runs:
    if ':' in r.text:
        r.text = f': {DATA["prodi"]}'

# [19] Sukabumi, ...
paras[19].text = f'Sukabumi, {DATA["tanggal"]}'

# [20] Staff Admin Akademik\tNama Mahasiswa
paras[20].text = f'Staff Admin Akademik\t\t\t\tMahasiswa'

# [24] (Tanda tangan baris bawah)
paras[24].text = f'( .................................................... )\t\t\t({DATA["nama"]})'

# --- 3. HALAMAN 3: LEMBAR PERSETUJUAN PELAKSANAAN SIDANG TESIS ---
# [32] Nama Mahasiswa :
for r in paras[32].runs:
    if ':' in r.text:
        r.text = f': {DATA["nama"]}'

# [34] NIM :
for r in paras[34].runs:
    if ':' in r.text:
        r.text = f': {DATA["nim"]}'

# [36] Program Studi :
for r in paras[36].runs:
    if ':' in r.text:
        r.text = f': {DATA["prodi"]}'

# [38] Alamat Rumah :
for r in paras[38].runs:
    if ':' in r.text:
        r.text = f': {DATA["alamat"]}'

# [40] Telepon (HP)/email :
# Clear existing runs and set precisely
for r in paras[40].runs:
    r.text = ''
r_label = paras[40].add_run('Telepon (HP)/email\t')
r_val = paras[40].add_run(f': {DATA["hp"]} / {DATA["email"]}')

# [42] Judul Penelitian :
for r in paras[42].runs:
    if ':' in r.text:
        r.text = f': {DATA["judul"]}'

# [51] Dosen Pembimbing Nama
paras[51].text = f'{DATA["ketua_pembimbing"]}\t{DATA["anggota_pembimbing"]}'
if paras[51].runs:
    paras[51].runs[0].bold = True

# [52] NIDN Pembimbing
paras[52].text = f'NIDN. {DATA["nidn_ketua_pembimbing"]}\tNIDN. {DATA["nidn_anggota_pembimbing"]}'

# [59] Ketua Program Studi ...
for r in paras[59].runs:
    if '...' in r.text:
        r.text = r.text.replace('...', DATA['prodi'])
if DATA['prodi'] not in paras[59].text:
    paras[59].text = f'Ketua Program Studi {DATA["prodi"]}'

# [65] Nama Kaprodi
paras[65].text = DATA['kaprodi']
if paras[65].runs:
    paras[65].runs[0].bold = True

# [66] NIDN Kaprodi
paras[66].text = f'NIDN : {DATA["nidn_kaprodi"]}'

# Save final updated docx
doc.save(TARGET_FILE)
print("SUCCESS: 026 FORM PERSYARATAN SIDANG TESIS.docx updated perfectly!")
