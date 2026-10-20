"""
Verification script for revised docx with UTF-8 stdout
"""

import docx
import re
import os
import sys

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

doc_path = "Revisi 4 A_Deep_Learning_Architecture_for_Cryptocurrency_Price_Forecasting.docx"
doc = docx.Document(doc_path)

print("=" * 60)
print("VERIFICATION AUDIT OF REVISED MANUSCRIPT")
print("=" * 60)

# Check paragraph counts and text
print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

# Check image paragraphs
images_found = 0
for i, p in enumerate(doc.paragraphs):
    blips = p._p.xpath('.//a:blip')
    if blips:
        images_found += len(blips)
        caption = doc.paragraphs[i+1].text if i+1 < len(doc.paragraphs) else "N/A"
        print(f"  [Image {images_found}] at P[{i}]: Caption -> {caption[:70]}")

print(f"Total Images in paragraphs: {images_found}")

# Check informal terms
informal_terms = ["fun part", "fix this", "You can't", "You can", "tells a very different story", "deploy to production"]
found_informal = False
for i, p in enumerate(doc.paragraphs):
    for t in informal_terms:
        if t.lower() in p.text.lower():
            print(f"  [WARNING] Informal term '{t}' at P[{i}]: {p.text[:80]}")
            found_informal = True
if not found_informal:
    print("  [OK] No informal or conversational language detected.")

# Check repetitions
p_texts = [p.text.strip() for p in doc.paragraphs if len(p.text.strip()) > 30]
duplicates = set([x for x in p_texts if p_texts.count(x) > 1])
if duplicates:
    print(f"  [WARNING] Duplicate paragraphs found: {len(duplicates)}")
    for d in duplicates:
        print(f"    Duplicate: {d[:60]}...")
else:
    print("  [OK] Zero duplicate paragraphs detected.")

# Check Table 0 (Algorithm 1)
print("\nAlgorithm 1 Content:")
for p in doc.tables[0].rows[0].cells[0].paragraphs[:10]:
    print("  ", p.text)

# Check Tables 1-3
for t_idx in [1, 2, 3]:
    print(f"\nTable {t_idx}:")
    for r in doc.tables[t_idx].rows[:2]:
        print("  ", " | ".join(c.text.strip() for c in r.cells))

print("\nVerification completed successfully!")
