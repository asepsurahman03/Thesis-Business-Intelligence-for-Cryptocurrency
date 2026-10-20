import docx
import re
import os
import latex2mathml.converter
from lxml import etree
import shutil

v7_path = '05_Reports/Manuscript_IJAI_FINAL_V7.docx'
v8_path = '05_Reports/Manuscript_IJAI_FINAL_V8.docx'

shutil.copy(v7_path, v8_path)
doc = docx.Document(v8_path)

xslt_path = r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"
xslt = etree.parse(xslt_path)
transform = etree.XSLT(xslt)

def latex_to_omml(latex_str):
    latex_str = latex_str.replace(r'\text{Attention}', r'\mathrm{Attention}')
    latex_str = latex_str.replace(r'\text{Softmax}', r'\mathrm{Softmax}')
    latex_str = latex_str.strip()
    try:
        mathml_str = latex2mathml.converter.convert(latex_str)
        dom = etree.fromstring(mathml_str)
        new_dom = transform(dom)
        return new_dom.getroot()
    except Exception as e:
        print("Error converting:", latex_str, e)
        return None

pattern = re.compile(r'(\$\$.*?\$\$|\$.*?\$)', re.DOTALL)

for p in doc.paragraphs:
    text = p.text
    if '$' in text:
        parts = pattern.split(text)
        
        p.clear()
        
        for part in parts:
            if part.startswith('$$') and part.endswith('$$'):
                math_str = part[2:-2].strip()
                omml = latex_to_omml(math_str)
                if omml is not None:
                    p._element.append(omml)
                else:
                    p.add_run(part)
            elif part.startswith('$') and part.endswith('$'):
                math_str = part[1:-1].strip()
                omml = latex_to_omml(math_str)
                if omml is not None:
                    p._element.append(omml)
                else:
                    p.add_run(part)
            else:
                if part:
                    run = p.add_run(part)
                    run.font.name = 'Times New Roman'

doc.save(v8_path)
print("Saved V8 with natively rendered OMML math.")
