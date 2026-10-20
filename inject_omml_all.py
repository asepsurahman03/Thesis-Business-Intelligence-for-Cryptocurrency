import docx
import re
import latex2mathml.converter
from lxml import etree
import shutil

input_path = '05_Reports/Manuscript_IJAI_FINAL_V5.docx'
output_path = '05_Reports/Manuscript_IJAI_FINAL_V5_FIXED.docx'

shutil.copy(input_path, output_path)
doc = docx.Document(output_path)

xslt_path = r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"
xslt = etree.parse(xslt_path)
transform = etree.XSLT(xslt)

def latex_to_omml(latex_str):
    latex_str = latex_str.replace(r'\text{Attention}', r'\mathrm{Attention}')
    latex_str = latex_str.replace(r'\text{Softmax}', r'\mathrm{Softmax}')
    latex_str = latex_str.strip()
    try:
        mathml_str = latex2mathml.converter.convert(latex_str)
        dom = etree.fromstring(mathml_str)
        new_dom = transform(dom)
        return new_dom.getroot()
    except Exception as e:
        print(f"Error converting: {latex_str} -> {e}")
        return None

def iter_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p

pattern = re.compile(r'(\$\$.*?\$\$|\$.*?\$)', re.DOTALL)

for p in iter_paragraphs(doc):
    text = p.text
    if '$' in text:
        parts = pattern.split(text)
        
        # We need to preserve alignment and style of paragraph if possible,
        # but p.clear() removes runs. It keeps paragraph format though!
        # Let's save the alignment and style if needed, but p.clear() only removes runs.
        # So paragraph formatting (like Center alignment) is preserved.
        p.clear()
        
        for part in parts:
            if part.startswith('$$') and part.endswith('$$'):
                math_str = part[2:-2].strip()
                omml = latex_to_omml(math_str)
                if omml is not None:
                    p._element.append(omml)
                else:
                    p.add_run(part)
            elif part.startswith('$') and part.endswith('$'):
                math_str = part[1:-1].strip()
                omml = latex_to_omml(math_str)
                if omml is not None:
                    p._element.append(omml)
                else:
                    p.add_run(part)
            else:
                if part:
                    run = p.add_run(part)
                    run.font.name = 'Times New Roman'

doc.save(output_path)
print("Saved FIXED with natively rendered OMML math across all paragraphs and tables.")
