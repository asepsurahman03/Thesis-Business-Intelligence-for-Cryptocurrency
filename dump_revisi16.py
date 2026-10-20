from docx import Document
doc = Document(r'c:\Nusa Putra University\S2 Magister\Thesis\Revisi_16_A_Deep_Learning_Architecture_for_Cryptocurrency_Price_Forecasting.docx')
with open(r'c:\Nusa Putra University\S2 Magister\Thesis\revisi16_dump.txt', 'w', encoding='utf-8') as f:
    for i, p in enumerate(doc.paragraphs):
        txt = ''.join(r.text for r in p.runs)
        if txt.strip():
            f.write(f'P{i}: {txt}\n\n')
    for ti, tbl in enumerate(doc.tables):
        f.write(f'=== TABLE {ti} ===\n')
        for ri, row in enumerate(tbl.rows):
            cells = [c.text.replace('\n',' ') for c in row.cells]
            f.write(f'Row {ri}: {chr(124).join(cells)}\n')
        f.write('\n')
print('Dump done.')
