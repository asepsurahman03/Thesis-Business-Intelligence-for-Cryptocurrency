import docx
import sys

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

doc = docx.Document('Revisi 4 A_Deep_Learning_Architecture_for_Cryptocurrency_Price_Forecasting.docx')

print("--- ALL CAPTIONS AND TABLE HEADINGS ---")
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if p.style.name in ['figure caption', 'table head', 'Caption'] or 'fig.' in txt.lower() or 'table ' in txt.lower():
        print(f"P[{i:03d}] (Style: {p.style.name}): '{txt}'")

print("\n--- ALL TABLES DETAIL ---")
for t_idx, t in enumerate(doc.tables):
    print(f"\nTable {t_idx} (rows: {len(t.rows)}, cols: {len(t.columns)}):")
    for r_idx, row in enumerate(t.rows):
        cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
        print(f"  Row {r_idx}: {' | '.join(cells)}")
