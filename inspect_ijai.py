import docx

doc = docx.Document('05_Reports/ijai.docx')
with open('ijai_structure.txt', 'w', encoding='utf-8') as f:
    f.write(f"Total Paragraphs: {len(doc.paragraphs)}\n\n")
    for i, p in enumerate(doc.paragraphs[:50]):
        f.write(f"[{i}] STYLE: {p.style.name} | TEXT: {p.text}\n")
    
    f.write("\n\n--- Source Manuscript ---\n\n")
    doc2 = docx.Document('05_Reports/Manuscript_IEEE_FINAL_2026_WORD.docx')
    for i, p in enumerate(doc2.paragraphs[:50]):
        f.write(f"[{i}] STYLE: {p.style.name} | TEXT: {p.text}\n")
