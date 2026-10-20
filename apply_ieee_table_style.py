"""
Apply exact IEEE/Scopus table styles, remove all duplicate caption prefixes,
and re-embed the updated 300 DPI figures without internal 'Fig. X' titles.
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
FIG_DIR = os.path.join(BASE_DIR, "04_Visualizations")
TARGET_DOC = os.path.join(BASE_DIR, "Revisi 4 A_Deep_Learning_Architecture_for_Cryptocurrency_Price_Forecasting.docx")
PDF_PATH = os.path.join(BASE_DIR, "Revisi 4 A_Deep_Learning_Architecture_for_Cryptocurrency_Price_Forecasting.pdf")

FIG1_PATH = os.path.join(FIG_DIR, "Fig1_BTC_Price_Volume_300dpi.png")
FIG2_PATH = os.path.join(FIG_DIR, "Fig2_Model_Comparison_300dpi.png")
FIG3_PATH = os.path.join(FIG_DIR, "Fig3_Residual_Distribution_300dpi.png")
FIG4_PATH = os.path.join(FIG_DIR, "Fig4_BI_Dashboard_KPI_300dpi.png")

doc = Document(TARGET_DOC)

def set_cell_borders_ieee(cell, is_header=False, is_last_row=False):
    """
    IEEE 3-line table format (Booktabs style):
    - Top of table has solid border (sz=8)
    - Bottom of header row has solid border (sz=4)
    - Bottom of table has solid border (sz=8)
    - No vertical lines anywhere (left=none, right=none, insideV=none)
    - No horizontal lines between data rows (insideH=none)
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    else:
        tcBorders.clear()

    # Define edges based on row position
    top_border = {'val': 'single', 'sz': '8', 'color': '000000', 'space': '0'} if is_header else {'val': 'none'}
    bottom_border = {'val': 'single', 'sz': '4', 'color': '000000', 'space': '0'} if is_header else (
        {'val': 'single', 'sz': '8', 'color': '000000', 'space': '0'} if is_last_row else {'val': 'none'}
    )

    borders = {
        'top': top_border,
        'bottom': bottom_border,
        'left': {'val': 'none'},
        'right': {'val': 'none'},
    }

    for edge, edge_data in borders.items():
        tag = f'w:{edge}'
        el = OxmlElement(tag)
        for k, v in edge_data.items():
            el.set(qn(f'w:{k}'), str(v))
        tcBorders.append(el)

def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def clear_cell_shading(cell):
    shd = cell._tc.get_or_add_tcPr().find(qn('w:shd'))
    if shd is not None:
        cell._tc.get_or_add_tcPr().remove(shd)

def replace_image_in_para(p, img_path, width_in_inches=3.4):
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(img_path, width=Inches(width_in_inches))

print("1. Updating Figures at 300 DPI...")
replace_image_in_para(doc.paragraphs[36], FIG1_PATH, width_in_inches=3.4)
replace_image_in_para(doc.paragraphs[64], FIG2_PATH, width_in_inches=3.4)
replace_image_in_para(doc.paragraphs[75], FIG3_PATH, width_in_inches=3.4)
replace_image_in_para(doc.paragraphs[94], FIG4_PATH, width_in_inches=3.4)

print("2. Ensuring clean figure captions without double prefixes...")
doc.paragraphs[37].text = "Fig. 1. Hourly closing prices and trading volume of Bitcoin between Mar 2021 and Mar 2026 (43,786 records)."
doc.paragraphs[37].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.paragraphs[65].text = "Fig. 2. 168-hour out-of-sample forecast overlay: actual BTC closing prices (solid) vs. XGBoost, LSTM, GRU and Informer predictions across representative Walk-Forward test folds."
doc.paragraphs[65].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.paragraphs[76].text = "Fig. 3. Residual error distributions for GRU and LSTM fold-aggregates: fat tails correspond to extreme ±3σ volatility episodes."
doc.paragraphs[76].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.paragraphs[95].text = "Fig. 4. Business Intelligence Dashboard KPI summary from 24-fold Walk-Forward evaluation (RMSE, Sharpe Ratio, Directional Accuracy, and VaR95%)."
doc.paragraphs[95].alignment = WD_ALIGN_PARAGRAPH.CENTER

print("3. Styling Table Titles in IEEE Scopus Standard...")
# Table I Title
doc.paragraphs[71].text = "TABLE I\nQUANTITATIVE LOSS BENCHMARKING ACROSS ALGORITHMIC WALK-FORWARD LAYERS"
doc.paragraphs[71].alignment = WD_ALIGN_PARAGRAPH.CENTER
for r in doc.paragraphs[71].runs:
    r.bold = True
    r.font.size = Pt(8.5)

# Table II Title
doc.paragraphs[83].text = "TABLE II\nKPI DASHBOARD (SHARPE, VAR, DIRECTIONAL ACCURACY)"
doc.paragraphs[83].alignment = WD_ALIGN_PARAGRAPH.CENTER
for r in doc.paragraphs[83].runs:
    r.bold = True
    r.font.size = Pt(8.5)

# Table III Title
doc.paragraphs[87].text = "TABLE III\nCOMPARATIVE BENCHMARKING AGAINST PUBLISHED LITERATURE"
doc.paragraphs[87].alignment = WD_ALIGN_PARAGRAPH.CENTER
for r in doc.paragraphs[87].runs:
    r.bold = True
    r.font.size = Pt(8.5)

print("4. Applying IEEE 3-line booktabs border format to Table 1, Table 2, Table 3...")
for t_idx in [1, 2, 3]:
    table = doc.tables[t_idx]
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    num_rows = len(table.rows)
    for r_idx, row in enumerate(table.rows):
        is_hdr = (r_idx == 0)
        is_last = (r_idx == num_rows - 1)
        for c_idx, cell in enumerate(row.cells):
            set_cell_borders_ieee(cell, is_header=is_hdr, is_last_row=is_last)
            set_cell_margins(cell, top=50, bottom=50, left=70, right=70)
            clear_cell_shading(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.05
                # Center header, left align first column, center data columns
                if is_hdr:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.bold = True
                        r.font.size = Pt(8.0)
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.size = Pt(7.5)

# Make sure Table 3 citation is [38] for Livieris
table3 = doc.tables[3]
for row in table3.rows:
    for cell in row.cells:
        if 'Livieris' in cell.text:
            cell.text = cell.text.replace('[40]', '[38]')
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(7.5)

# Re-save the document
doc.save(TARGET_DOC)
print(f"[OK] Successfully saved updated document: {TARGET_DOC}")

# Re-export PDF
print("5. Exporting updated PDF via Word COM...")
try:
    import win32com.client
    word = win32com.client.Dispatch('Word.Application')
    word.Visible = False
    doc_com = word.Documents.Open(os.path.abspath(TARGET_DOC))
    doc_com.SaveAs(os.path.abspath(PDF_PATH), FileFormat=17)
    doc_com.Close()
    word.Quit()
    print(f"[OK] PDF successfully exported to: {PDF_PATH}")
except Exception as e:
    print(f"[!] PDF export warning: {e}")
