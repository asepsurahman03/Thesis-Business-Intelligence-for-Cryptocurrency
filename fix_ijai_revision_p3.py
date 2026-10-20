"""
IJ-AI Manuscript Revision - Phase 3 Final
UTF-8 output. Fixes remaining items.
"""

import sys
import re
import io
from docx import Document

# Force UTF-8 output
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

SRC = r"e:\Nusa Putra\S2\Tesis\Analisis Bit Coin\05_Reports\Deep_Learning_Cryptocurrency_BI_REVISED_V2.docx"
OUT = r"e:\Nusa Putra\S2\Tesis\Analisis Bit Coin\05_Reports\Deep_Learning_Cryptocurrency_BI_REVISED_V3.docx"

doc = Document(SRC)

def para_text(p):
    return "".join(r.text for r in p.runs)

def set_text(p, new_text):
    if p.runs:
        for r in p.runs:
            r.text = ""
        p.runs[0].text = new_text
    else:
        p.add_run(new_text)

print("=== PHASE 3 REVISION ===\n")

# ─────────────────────────────────────────────────────────────────────────────
# FIX A: Keywords paragraph — scan all early paras
# Diagnostic showed no keyword para in early paras (empty keyword scan)
# The keyword line is likely in a table cell or has specific XML styling
# ─────────────────────────────────────────────────────────────────────────────
print("[FIX A] Finding and fixing Keywords...")
# Check table cells too (keywords are sometimes in a table)
kw_fixed = False

# First check regular paragraphs
for i, p in enumerate(doc.paragraphs):
    txt = para_text(p)
    if i > 20:
        break
    if txt.strip():
        print(f"  Para[{i}]: {txt[:80]}")

# Also search tables
print("\n  Searching tables for keywords...")
for t_idx, table in enumerate(doc.tables):
    for row in table.rows:
        for cell in row.cells:
            cell_txt = cell.text
            if ("LSTM" in cell_txt and "GRU" in cell_txt and 
                ("Business Intelligence" in cell_txt or "Keywords" in cell_txt.split("\n")[0])):
                print(f"  Found in table {t_idx}: {cell_txt[:150]}")
                # Fix in cell paragraphs
                for cp in cell.paragraphs:
                    cp_txt = para_text(cp)
                    if "LSTM" in cp_txt or "GRU" in cp_txt:
                        new_txt = re.sub(r'\bLSTM\b', 'Long Short-Term Memory (LSTM)', cp_txt)
                        new_txt = re.sub(r'\bGRU\b', 'Gated Recurrent Unit (GRU)', new_txt)
                        if new_txt != cp_txt:
                            set_text(cp, new_txt)
                            print(f"  Keywords updated in table {t_idx}")
                            kw_fixed = True

if not kw_fixed:
    print("  Scanning all paragraphs for keyword line...")
    for i, p in enumerate(doc.paragraphs):
        txt = para_text(p)
        style = p.style.name if p.style else ""
        # Keywords line usually has style "Keywords" or is short list with commas and | separator
        if (("Keywords" in txt[:15] or style == "Keywords" or 
             ("Financial Informatics" in txt and "LSTM" in txt)) and 
            ("LSTM" in txt or "GRU" in txt)):
            new_txt = re.sub(r'\bLSTM\b', 'Long Short-Term Memory (LSTM)', txt)
            new_txt = re.sub(r'\bGRU\b', 'Gated Recurrent Unit (GRU)', new_txt)
            if new_txt != txt:
                set_text(p, new_txt)
                print(f"  Keywords fixed at para [{i}]")
                kw_fixed = True
                break

if not kw_fixed:
    print("  [NOTE] Keywords not found in paragraphs - may be in a text box or header")
    print("  [ACTION] Open .docx manually and change: LSTM -> Long Short-Term Memory (LSTM), GRU -> Gated Recurrent Unit (GRU)")

# ─────────────────────────────────────────────────────────────────────────────
# FIX B: [23] TimesNet citation
# Para [21] contains Informer + PatchTST but already has [23]? 
# Let's check carefully
# ─────────────────────────────────────────────────────────────────────────────
print("\n[FIX B] Adding [23] TimesNet citation...")
p21 = doc.paragraphs[21]
p21_txt = para_text(p21)
print(f"  Para[21] (first 300 chars): {p21_txt[:300]}")

if "TimesNet" in p21_txt:
    if "[23]" not in p21_txt:
        new_txt = p21_txt.replace("TimesNet", "TimesNet [23]", 1)
        set_text(p21, new_txt)
        print("  Added [23] after TimesNet at para [21]")
    else:
        print("  [23] already cited after TimesNet - OK")
else:
    # Search entire doc
    for i, p in enumerate(doc.paragraphs):
        txt = para_text(p)
        if "PatchTST" in txt and "[22]" in txt and "[23]" not in txt:
            new_txt = txt.replace("[22]", "[22], [23]", 1)
            set_text(p, new_txt)
            print(f"  Added [23] citation at para [{i}]")
            break
        elif "TimesNet" in txt and "[23]" not in txt:
            new_txt = txt.replace("TimesNet", "TimesNet [23]", 1)
            set_text(p, new_txt)
            print(f"  Added [23] citation at para [{i}]")
            break

# ─────────────────────────────────────────────────────────────────────────────
# FIX C: Table III pre-citation — check Results area
# From diagnostic: "Table 1." seen at [117,120,122] - the doc uses "Table 1" not "Table I"
# ─────────────────────────────────────────────────────────────────────────────
print("\n[FIX C] Checking Table citations (format 'Table 1.' vs 'Table I')...")
# The doc appears to use numbered format "Table 1." etc.
# Check if Table 3 (KPI) has pre-citation
kpi_table_found = False
for i, p in enumerate(doc.paragraphs):
    txt = para_text(p).strip()
    if ("Table 3" in txt or "Table III" in txt) and i < 200:
        print(f"  Table 3/III found at [{i}]: {txt[:100]}")
        kpi_table_found = True

if not kpi_table_found:
    print("  [!] Table 3/III not found - searching for KPI table...")
    for i, p in enumerate(doc.paragraphs):
        txt = para_text(p)
        if "KPI" in txt and ("Sharpe" in txt or "VaR" in txt) and i > 100:
            print(f"  KPI content at [{i}]: {txt[:100]}")

# Add Table 3 pre-citation if missing
for i, p in enumerate(doc.paragraphs):
    txt = para_text(p).strip()
    if ("Fulfilling the targeted architectural bridge" in txt or
        "Table III outlines" in txt or
        "KPI metrics evaluating models specifically" in txt):
        if "Table" not in txt:
            new_txt = txt.rstrip(".")
            new_txt += (
                ". Table 3 presents the translated KPI outputs—including Directional Accuracy, "
                "Sharpe Ratio, and VaR(95%)—comparing the four evaluated model architectures."
            )
            set_text(p, new_txt)
            print(f"  Added Table 3 pre-citation at [{i}]")
        else:
            print(f"  Table pre-cited at [{i}] - OK")
        break

# ─────────────────────────────────────────────────────────────────────────────
# FIX D: Verify Fig.1 is cited before the caption (para 78)
# From diagnostic: captions start with "Figure 1." (not "Fig. 1.")
# ─────────────────────────────────────────────────────────────────────────────
print("\n[FIX D] Checking Fig. 1 pre-citation (caption is at para 78)...")
# Look for any paragraph before 78 that references Figure 1
fig1_precited = False
for i in range(0, 78):
    txt = para_text(doc.paragraphs[i])
    if "Figure 1" in txt or "Fig. 1" in txt or "Fig.1" in txt:
        print(f"  Fig.1 pre-cited at para [{i}] - OK")
        fig1_precited = True
        break

if not fig1_precited:
    print("  [!] Fig. 1 not pre-cited before para 78 - searching for 43,786 paragraph...")
    for i in range(0, 78):
        txt = para_text(doc.paragraphs[i]).strip()
        if "43,786" in txt and "Figure 1" not in txt:
            new_txt = txt.rstrip(".")
            new_txt += (
                ". Figure 1 illustrates the full BTC/USDT price and volume history "
                "across this 5-year observation window."
            )
            set_text(doc.paragraphs[i], new_txt)
            print(f"  Added Fig.1 pre-citation at para [{i}]")
            fig1_precited = True
            break

# ─────────────────────────────────────────────────────────────────────────────
# FIX E: Missing DOIs - check refs starting at ~165
# ─────────────────────────────────────────────────────────────────────────────
print("\n[FIX E] Checking references for missing DOIs...")
ref_dois = {
    "F. A. Gers, J. Schmidhuber, and F. Cummins": 
        " doi: 10.1162/089976600300015015.",
    "K. Greff et al":
        " doi: 10.1109/TNNLS.2016.2582924.",
    "A. Vaswani et al":
        " doi: 10.5555/3295222.3295349.",
    "R. S. Tsay":
        " doi: 10.1002/9780470644560.",
    "S. M. Lundberg and S. I. Lee":
        " doi: 10.5555/3295222.3295230.",
}

for i, p in enumerate(doc.paragraphs):
    txt = para_text(p)
    for ref_key, doi_suffix in ref_dois.items():
        if ref_key in txt and "doi:" not in txt.lower() and "10." not in txt:
            new_txt = txt.rstrip(".") + doi_suffix
            set_text(p, new_txt)
            print(f"  Added DOI for '{ref_key[:40]}...' at para [{i}]")
            break

# ─────────────────────────────────────────────────────────────────────────────
# FIX F: Fix para at [13] and [14] in intro - verify expansion happened
# ─────────────────────────────────────────────────────────────────────────────
print("\n[FIX F] Verifying intro paragraphs expanded...")
for i in [13, 14, 25]:
    if i < len(doc.paragraphs):
        txt = para_text(doc.paragraphs[i])
        count = len([s.strip() for s in txt.split('.') if s.strip()])
        status = "OK (3+ sentences)" if count >= 3 else f"[!] Only {count} sentence(s)"
        print(f"  Para [{i}]: {status} | {txt[:80]}")

# ─────────────────────────────────────────────────────────────────────────────
# FIX G: Ensure 3-sentence rule for key sections
# Section 3 opening, 4 opening, 5 opening - verify
# ─────────────────────────────────────────────────────────────────────────────
print("\n[FIX G] Verifying other key paragraphs have 3+ sentences...")
important_checks = [
    "four-layer architecture mapping data ingestions",
    "univariate sequences ignores market momentum",
    "720-hour (30-day) training step paired",
    "empirical outperformance of the GRU",
    "Livieris et al",
]
for check in important_checks:
    for i, p in enumerate(doc.paragraphs):
        txt = para_text(p)
        if check in txt:
            count = len([s.strip() for s in txt.split('.') if s.strip()])
            status = "OK" if count >= 3 else f"[!] {count} sentence(s)"
            print(f"  '{check[:35]}...' [{i}]: {status} ({len(txt)} chars)")
            break

# ─────────────────────────────────────────────────────────────────────────────
# SAVE
# ─────────────────────────────────────────────────────────────────────────────
print("\n" + "=" * 60)
doc.save(OUT)
print(f"[OK] Saved: {OUT}")
print("=" * 60)
