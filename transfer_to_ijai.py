import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
import copy

def copy_paragraph(p, doc):
    new_p = doc.add_paragraph()
    # Try to copy style if it exists in the target document, otherwise use Normal
    try:
        new_p.style = p.style.name
    except KeyError:
        new_p.style = 'Normal'
    
    new_p.alignment = p.alignment
    
    for run in p.runs:
        new_run = new_p.add_run(run.text)
        new_run.bold = run.bold
        new_run.italic = run.italic
        new_run.underline = run.underline
        if run.font.size:
            new_run.font.size = run.font.size
        if run.font.name:
            new_run.font.name = run.font.name

source_doc = docx.Document('05_Reports/Manuscript_IEEE_FINAL_2026_WORD.docx')
template_doc = docx.Document('05_Reports/ijai.docx')

# Clear the template document paragraphs to start fresh but keep its styles
for p in template_doc.paragraphs:
    p._element.getparent().remove(p._element)

# Clear tables as well
for table in template_doc.tables:
    table._element.getparent().remove(table._element)

# Iterate over all elements (paragraphs and tables) in the source doc's body
# python-docx doesn't provide a direct way to iterate paragraphs and tables in order easily without dropping to lxml
# But we can approximate by reading the body elements directly
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

body = source_doc._body._body

for child in body.iterchildren():
    if isinstance(child, CT_P):
        p = Paragraph(child, source_doc)
        copy_paragraph(p, template_doc)
    elif isinstance(child, CT_Tbl):
        table = Table(child, source_doc)
        new_table = template_doc.add_table(rows=len(table.rows), cols=len(table.columns))
        new_table.style = 'Table Grid'
        
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                # Only copy text for simplicity, merging is complex
                try:
                    new_table.cell(r_idx, c_idx).text = cell.text
                except IndexError:
                    pass

template_doc.save('05_Reports/Manuscript_IJAI_FINAL.docx')
print("Transfer complete. Saved to 05_Reports/Manuscript_IJAI_FINAL.docx")
