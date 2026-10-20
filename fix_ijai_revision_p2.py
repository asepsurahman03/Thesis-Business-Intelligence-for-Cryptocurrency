"""
IJ-AI Manuscript Revision Script - Phase 2
Fixes remaining issues found after Phase 1:
1. Keywords (LSTM/GRU full forms) - re-examine
2. 3 thin intro paragraphs
3. [23] TimesNet citation not inserted
4. Figure captions not updated (search key mismatch)
5. Table III pre-citation check
6. Table I / Fig 1 pre-citation
7. Diebold-Mariano reference
"""

import re
from docx import Document

SRC = r"e:\Nusa Putra\S2\Tesis\Analisis Bit Coin\05_Reports\Deep_Learning_Cryptocurrency_BI_REVISED.docx"
OUT = r"e:\Nusa Putra\S2\Tesis\Analisis Bit Coin\05_Reports\Deep_Learning_Cryptocurrency_BI_REVISED_V2.docx"

doc = Document(SRC)

def para_text(p):
    return "".join(r.text for r in p.runs)

def set_text(p, new_text):
    if p.runs:
        for r in p.runs:
            r.text = ""
        p.runs[0].text = new_text
    else:
        p.add_run(new_text)

# ─────────────────────────────────────────────────────────────────────────────
# DIAGNOSTIC: Dump all paragraphs containing keywords we care about
# ─────────────────────────────────────────────────────────────────────────────
print("=== DIAGNOSTIC SCAN ===")
keywords_to_find = [
    "Keywords", "LSTM", "GRU", "PatchTST", "TimesNet",
    "Fig. 1", "Fig. 2", "Fig. 3", "Fig. 4",
    "Table I", "Table II", "Table III",
    "Hourly closing prices",
    "Performance comparison mapping",
    "Statistical configuration distribution",
    "Web-based prototype",
    "Deden Witarsyah",
    "neural networks have become preferred",
    "critical gap exists between",
    "BTC/USDT was selected as the sole asset",
]

for kw in keywords_to_find:
    for i, p in enumerate(doc.paragraphs):
        txt = para_text(p)
        if kw.lower() in txt.lower():
            snippet = txt[:120].replace('\n', ' ')
            print(f"  [{i:3d}] '{kw}': {snippet!r}")
            break

print("\n")

# ─────────────────────────────────────────────────────────────────────────────
# FIX A: Keywords — find correct paragraph and expand
# ─────────────────────────────────────────────────────────────────────────────
print("[FIX A] Expanding keywords LSTM -> full form, GRU -> full form...")
fixed_kw = False
for i, p in enumerate(doc.paragraphs):
    txt = para_text(p)
    if ("Keywords" in txt or "keywords" in txt) and len(txt) > 30:
        if "LSTM" in txt or "GRU" in txt:
            new_txt = txt
            new_txt = re.sub(r'\bLSTM\b', 'Long Short-Term Memory (LSTM)', new_txt)
            new_txt = re.sub(r'\bGRU\b', 'Gated Recurrent Unit (GRU)', new_txt)
            set_text(p, new_txt)
            print(f"  Keywords updated at paragraph {i}")
            print(f"  New text: {new_txt[:150]!r}")
            fixed_kw = True
            break
if not fixed_kw:
    print("  [!] Keywords paragraph NOT found - searching all paragraphs with LSTM/GRU...")
    for i, p in enumerate(doc.paragraphs):
        txt = para_text(p)
        if "Business Intelligence" in txt and "LSTM" in txt and "GRU" in txt and len(txt) < 300:
            new_txt = re.sub(r'\bLSTM\b', 'Long Short-Term Memory (LSTM)', txt)
            new_txt = re.sub(r'\bGRU\b', 'Gated Recurrent Unit (GRU)', new_txt)
            set_text(p, new_txt)
            print(f"  Keywords updated at paragraph {i} (fallback)")
            break

# ─────────────────────────────────────────────────────────────────────────────
# FIX B: Expand 3 thin intro paragraphs (found at 13, 14, 25)
# ─────────────────────────────────────────────────────────────────────────────
print("\n[FIX B] Expanding thin introduction paragraphs...")

thin_para_expansions = {
    "Data-driven neural networks have become preferred for modeling complex non-linea": (
        "Data-driven neural networks have become preferred for modeling complex non-linear "
        "financial sequences absent in classical statistics [5], [6]. "
        "However, evaluating advanced architectures like Recurrent Neural Networks (RNN) or "
        "state-of-the-art Transformers strictly upon standard error metrics (e.g., RMSE) is "
        "insufficient for algorithmic portfolio management, as it neglects the operational "
        "dimensions of trading profitability [7], [8]. "
        "Effective deployment requires translating model outputs into actionable risk "
        "indicators that are interpretable by institutional decision-makers."
    ),
    "A critical gap exists between algorithmic accuracy and operational deployment wi": (
        "A critical gap exists between theoretical algorithmic accuracy and operational "
        "deployment within enterprise Business Intelligence (BI) platforms [9]. "
        "Forecasting frameworks require an architecture capable of bridging raw numerical "
        "outputs into quantifiable risk metrics and Key Performance Indicators (KPI) "
        "compatible with institutional risk management standards [10]. "
        "Without this integration layer, even highly accurate models remain disconnected "
        "from the decision workflows of financial practitioners."
    ),
    "BTC/USDT was selected as the sole asset because Bitcoin maintains the highest ma": (
        "BTC/USDT was selected as the sole asset because Bitcoin maintains the highest market "
        "capitalization, trading volume, and liquidity among all cryptocurrencies [1], [30]. "
        "Its continuous 24/7 traded nature at hourly granularity provides the highest-density "
        "sequential signal environment for temporal model benchmarking, free from the "
        "liquidity discontinuities affecting smaller-cap assets. "
        "Restricting the analysis to a single, institutionally dominant asset also functions "
        "as an intentional experimental control, ensuring that measured framework performance "
        "reflects algorithmic and BI pipeline quality rather than cross-asset microstructural variance."
    ),
}

for search_prefix, expanded_text in thin_para_expansions.items():
    found = False
    for i, p in enumerate(doc.paragraphs):
        txt = para_text(p).strip()
        if txt.startswith(search_prefix[:60]):
            set_text(p, expanded_text)
            print(f"  Expanded intro paragraph at {i}: '{search_prefix[:50]}...'")
            found = True
            break
    if not found:
        print(f"  [!] NOT FOUND: '{search_prefix[:50]}...'")

# ─────────────────────────────────────────────────────────────────────────────
# FIX C: Add [23] TimesNet citation — comprehensive search
# ─────────────────────────────────────────────────────────────────────────────
print("\n[FIX C] Adding [23] TimesNet citation in text...")
timesnet_added = 0
for i, p in enumerate(doc.paragraphs):
    txt = para_text(p)
    if "PatchTST" in txt and "[22]" in txt and "[23]" not in txt:
        # Try multiple replacement patterns
        new_txt = txt.replace("PatchTST [22] and iTransformer represent",
                               "PatchTST [22], TimesNet [23], and iTransformer represent")
        if new_txt == txt:
            new_txt = txt.replace("PatchTST [22]", "PatchTST [22] and TimesNet [23]")
        if new_txt == txt:
            # last resort: insert after [22]
            new_txt = txt.replace("[22]", "[22], [23]", 1)
        if new_txt != txt:
            set_text(p, new_txt)
            print(f"  [23] inserted at paragraph {i}")
            timesnet_added += 1
        else:
            print(f"  [!] Could not insert [23] at paragraph {i}: {txt[:100]!r}")

if timesnet_added == 0:
    print("  [!] PatchTST paragraph not found - checking all paras with iTransformer...")
    for i, p in enumerate(doc.paragraphs):
        txt = para_text(p)
        if "iTransformer" in txt and "[22]" in txt:
            new_txt = txt.replace("[22]", "[22], [23]", 1)
            set_text(p, new_txt)
            print(f"  [23] inserted at paragraph {i} (fallback): {new_txt[:80]!r}")
            timesnet_added += 1
            break

# ─────────────────────────────────────────────────────────────────────────────
# FIX D: Figure captions — comprehensive search & update
# ─────────────────────────────────────────────────────────────────────────────
print("\n[FIX D] Updating figure captions with axis labels...")

fig_updates = [
    {
        "search": "Hourly closing prices",
        "new": ("Fig. 1. Hourly BTC/USDT closing prices (left Y-axis: Closing Price in USDT) and "
                "trading volume (right Y-axis: Volume in BTC) between Mar 2021 and Mar 2026 "
                "(43,786 hourly records). X-axis: Date timeline at 1-hour granularity.")
    },
    {
        "search": "Performance comparison mapping localized 7-day",
        "new": ("Fig. 2. Performance comparison mapping localized 7-day (168-hour) test trajectory "
                "fluctuations between standard regressions and sequence networks. "
                "X-axis: Walk-Forward fold time step (hours 1–168); "
                "Y-axis: Normalized closing price (Min-Max scaled, 0–1).")
    },
    {
        "search": "Statistical",
        "neg_check": "Hourly",  # exclude Fig.1 false match
        "new": ("Fig. 3. Statistical residual error distribution illustrating localized error clusters "
                "surrounding the GRU baseline against extreme +3sigma tail events. "
                "X-axis: Normalized residual error magnitude; "
                "Y-axis: Frequency count across all 26 Walk-Forward folds.")
    },
    {
        "search": "Web-based prototype implementation",
        "new": ("Fig. 4. Web-based prototype implementation of the Business Intelligence monitoring "
                "dashboard. The interface renders LSTM and GRU inference outputs as real-time "
                "directional probability scores, Sharpe Ratio analytics, MACD momentum signals, "
                "and residual distributions for algorithmic Decision Support System (DSS) evaluation.")
    },
]

for fig in fig_updates:
    search_key = fig["search"]
    new_cap = fig["new"]
    neg = fig.get("neg_check", None)
    
    for i, p in enumerate(doc.paragraphs):
        txt = para_text(p)
        if search_key in txt:
            if neg and neg in txt:
                continue
            # Check it looks like a caption (starts with Fig or is near a fig)
            if txt.strip().startswith("Fig") or "Fig." in txt[:10]:
                set_text(p, new_cap)
                print(f"  Updated Fig caption at {i}: '{search_key}'")
                break
            else:
                # Some docx files store caption in a special paragraph style
                style_name = p.style.name if p.style else "?"
                print(f"  Found '{search_key}' at {i} (style: {style_name}): {txt[:80]!r}")
                if "Caption" in style_name or "caption" in style_name:
                    set_text(p, new_cap)
                    print(f"    -> Updated (Caption style)")
                    break

# ─────────────────────────────────────────────────────────────────────────────
# FIX E: Ensure Fig. 1 is cited before it appears (in Section 3.1)
# ─────────────────────────────────────────────────────────────────────────────
print("\n[FIX E] Ensuring Fig. 1 pre-citation...")
for i, p in enumerate(doc.paragraphs):
    txt = para_text(p).strip()
    if "43,786 historical OHLCV records" in txt and "Figure 1" not in txt and "Fig. 1" not in txt:
        new_txt = txt.rstrip(".")
        new_txt += (
            ". Figure 1 illustrates the full temporal coverage of the BTC/USDT closing "
            "price series and corresponding trading volume across the 5-year observation window."
        )
        set_text(p, new_txt)
        print(f"  Added Fig.1 pre-citation at paragraph {i}")
        break

# ─────────────────────────────────────────────────────────────────────────────
# FIX F: Ensure Table III is cited before table appears
# ─────────────────────────────────────────────────────────────────────────────
print("\n[FIX F] Ensuring Table III pre-citation...")
for i, p in enumerate(doc.paragraphs):
    txt = para_text(p).strip()
    if "Fulfilling the targeted architectural bridge" in txt and "Table III" not in txt:
        new_txt = txt.replace(
            "Fulfilling the targeted architectural bridge detailed in Algorithm 1,",
            "Fulfilling the targeted architectural bridge detailed in Algorithm 1, Table III"
        )
        if new_txt == txt:
            new_txt = txt.rstrip(".")
            new_txt += " Table III presents the KPI dashboard outputs for all four models."
        set_text(p, new_txt)
        print(f"  Updated Table III pre-citation at paragraph {i}")
        break
else:
    # Already has Table III citation
    for i, p in enumerate(doc.paragraphs):
        txt = para_text(p).strip()
        if "Table III outlines" in txt or ("architectural bridge" in txt and "Table III" in txt):
            print(f"  Table III pre-citation OK at paragraph {i}")
            break

# ─────────────────────────────────────────────────────────────────────────────
# FIX G: Add Diebold-Mariano Ref [36] DOI if missing
# ─────────────────────────────────────────────────────────────────────────────
print("\n[FIX G] Verifying Diebold-Mariano DOI in references...")
for i, p in enumerate(doc.paragraphs):
    txt = para_text(p)
    if "Diebold and R. Mariano" in txt:
        if "doi:" not in txt.lower():
            new_txt = txt.rstrip(".") + " doi: 10.1080/07350015.1995.10524599."
            set_text(p, new_txt)
            print(f"  Added DOI for Diebold-Mariano at paragraph {i}")
        else:
            print(f"  Diebold-Mariano DOI already present at paragraph {i}")
        break

# ─────────────────────────────────────────────────────────────────────────────
# FIX H: Fix paragraph that starts "The proposed framework" (Section 3 opening)
# making sure it was expanded in Phase 1
# ─────────────────────────────────────────────────────────────────────────────
print("\n[FIX H] Verifying Section 3 opening paragraph expansion...")
for i, p in enumerate(doc.paragraphs):
    txt = para_text(p).strip()
    if "four-layer architecture mapping data ingestions" in txt:
        if len(txt) < 100:
            # Still single sentence - expand it
            new_txt = (
                "The proposed framework adopts a four-layer architecture mapping data ingestions "
                "directly into a BI interface. This pipeline integrates raw Binance API feeds "
                "through ETL transformation, Walk-Forward model inference, and a PostgreSQL Data Mart "
                "into a Microsoft Power BI dashboard for real-time risk metric computation. "
                "Each layer is designed to be modular and replaceable, enabling future extensions "
                "such as multi-asset feeds or alternative deep learning architectures without "
                "redesigning the entire pipeline."
            )
            set_text(p, new_txt)
            print(f"  Expanded Section 3 opening at paragraph {i}")
        else:
            print(f"  Section 3 opening already expanded at paragraph {i} ({len(txt)} chars)")
        break

# ─────────────────────────────────────────────────────────────────────────────
# SAVE
# ─────────────────────────────────────────────────────────────────────────────
print("\n" + "=" * 60)
print("Saving Phase 2 revised manuscript...")
doc.save(OUT)
print(f"[OK] Saved: {OUT}")
print("=" * 60)
print("\nRevision Phase 2 complete.")
