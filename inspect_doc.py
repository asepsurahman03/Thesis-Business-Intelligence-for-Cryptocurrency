import docx
import os
import re

doc_path = "Revisi 4 A_Deep_Learning_Architecture_for_Cryptocurrency_Price_Forecasting.docx"
doc = docx.Document(doc_path)

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

with open("doc_full_text.txt", "w", encoding="utf-8") as out:
    for i, p in enumerate(doc.paragraphs):
        out.write(f"=== PARAGRAPH {i} (Style: {p.style.name}) ===\n{p.text}\n\n")
    
    for t_idx, table in enumerate(doc.tables):
        out.write(f"\n=== TABLE {t_idx} ===\n")
        for r_idx, row in enumerate(table.rows):
            cells = [c.text.strip().replace('\n', ' [NL] ') for c in row.cells]
            out.write(f"Row {r_idx}: {' | '.join(cells)}\n")

print("Full text saved to doc_full_text.txt")
