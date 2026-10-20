import docx

doc = docx.Document('05_Reports/Manuscript_IJAI_FINAL_V9.docx')

print(f"Total tables: {len(doc.tables)}")
for i, table in enumerate(doc.tables):
    print(f"\nTable {i+1} dimensions: {len(table.rows)} rows, {len(table.columns)} columns")
    print("Table Alignment:", table.alignment)
    print("Table Style:", table.style.name if table.style else "None")
    
    # Print the first few cells to identify the table
    try:
        header = [cell.text.replace('\n', ' ').strip() for cell in table.rows[0].cells]
        print("Header:", header)
    except:
        print("Could not read header.")
        
    # Check paragraph before the table (usually the caption)
    # Finding table position is tricky in python-docx, but we can iterate block-level elements
