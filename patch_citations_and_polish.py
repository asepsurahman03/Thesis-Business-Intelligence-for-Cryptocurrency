"""
Patch missing citations and polish document:
Integrates [17], [25], [33], [35] seamlessly into body text.
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Pt

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TARGET_DOC = os.path.join(BASE_DIR, "Revisi 4 A_Deep_Learning_Architecture_for_Cryptocurrency_Price_Forecasting.docx")
PDF_PATH = os.path.join(BASE_DIR, "Revisi 4 A_Deep_Learning_Architecture_for_Cryptocurrency_Price_Forecasting.pdf")

doc = Document(TARGET_DOC)

# 1. Patch Ref [17]: Gers et al. at P[19]
p19 = doc.paragraphs[19]
p19.text = (
    "To capture multi-scale temporal dependencies, recurrent neural architectures were introduced. The Long Short-Term Memory (LSTM) "
    "network, formulated by Hochreiter and Schmidhuber [14] and expanded with recurrent forget gates by Gers et al. [17], regulates "
    "gradient flow via input, forget, and output gating mechanisms, mitigating vanishing gradient phenomena. Subsequently, Cho et al. [15] "
    "developed the Gated Recurrent Unit (GRU), which combines forget and input gates into a streamlined update gate, achieving comparable "
    "sequential retention with reduced computational overhead. Both topologies remain essential benchmarks in quantitative financial engineering [16], [18]."
)

# 2. Patch Ref [25]: Baqee et al. at P[21]
p21 = doc.paragraphs[21]
p21.text = (
    "Attention-based Transformer architectures [19] have achieved remarkable success across diverse sequential modeling domains [20]. "
    "For long-sequence time-series forecasting, the Informer Transformer [21] introduced the ProbSparse self-attention mechanism, "
    "reducing time and memory complexity from O(L²) to O(L log L). While newer Transformer derivatives such as PatchTST [22] and "
    "TimesNet [23] continue to emerge, the Informer serves as an established, extensively validated attention baseline for intra-day "
    "cryptocurrency modeling [25]. Importantly, Zerveas et al. [24] observed that attention mechanisms are prone to overfitting on short, "
    "noisy intervals—a phenomenon directly investigated in this study."
)

# 3. Patch Ref [33]: Tsay at P[32]
p32 = doc.paragraphs[32]
p32.text = (
    "To guarantee zero temporal leakage across the chronological evaluation, all feature normalization was performed using strictly "
    "Fold-Local Min-Max Scaling (Equation 1). Scaling parameters (Xmin, Xmax) were computed exclusively from historical observations "
    "within each training window and subsequently applied out-of-sample to the corresponding test window, precluding any look-ahead bias [31], [32], [33]:"
)

# 4. Patch Ref [35]: Lundberg & Lee (SHAP) at P[38] / P[39]
p39 = doc.paragraphs[39]
p39.text = (
    "The 15 technical features incorporate trend indicators (EMA12, EMA26), momentum oscillators (RSI, Stochastic %K/%D), "
    "volatility bands (Bollinger Bands, Average True Range), and moving average convergence. Furthermore, SHapley Additive exPlanations (SHAP) "
    "analysis [35] was conducted across all models, confirming that volatility and momentum measures contributed the dominant predictive attribution:"
)

doc.save(TARGET_DOC)
print(f"[OK] Patched citations in: {TARGET_DOC}")

# Re-export PDF
try:
    import win32com.client
    word = win32com.client.Dispatch('Word.Application')
    word.Visible = False
    doc_com = word.Documents.Open(os.path.abspath(TARGET_DOC))
    doc_com.SaveAs(os.path.abspath(PDF_PATH), FileFormat=17)
    doc_com.Close()
    word.Quit()
    print(f"[OK] PDF re-exported to: {PDF_PATH}")
except Exception as e:
    print(f"[!] PDF export error: {e}")
